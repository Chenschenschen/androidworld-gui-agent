"""Per-run log cards for the submission (course log fields, D1 p.49, plus note, trajectory link and agent version).

usage: python build_run_log.py <repo>   -> <repo>/docs/RUN_LOG.md and Run_Log_Final_Evaluation.docx next to this file
One table per task: fields as rows, the three runs as columns (the layout our team used for its FilesMoveFile comparison).
"""
import sys
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from docx.enum.section import WD_ORIENT
from run_log_data import RUNS, ORDER, ROW, OBS, BUDGET, MODEL, VERSION, ZIP, trajectory_url

HERE = Path(__file__).resolve().parent
repo = Path(sys.argv[1])
clips = {r[3]: sorted(v.name for v in (repo / "runs" / r[3]).glob("video_*.mp4")) for r in RUNS}
FIELDS = ["task", "run", "model", "observation", "grounding", "steps", "verifier", "failure_class", "recording", "note",
          "trajectory", "agent version"]


def values(run):
    task, number, seed, run_id, steps, verdict, failure, n_clips, short, note = run
    return {"task": task, "run": f"{number} of 3 (seed {seed})", "model": MODEL, "observation": OBS.get(task, "screenshot"),
            "grounding": "coordinates (0-1000 grid)", "steps": f"{steps} (budget {BUDGET[task]})", "verifier": verdict,
            "failure_class": failure or "—", "recording": ", ".join(clips[run_id]) + f" in {ZIP} / {run_id}",
            "note": note, "trajectory": run_id, "agent version": VERSION}


# Markdown (renders on GitHub, links work there)
md = ["# Run log: final evaluation (v65)\n",
      "One card per task; the three runs are the columns. Fields follow the course's run log (D1 p.49), plus a note that explains "
      "the run, a link to its readable trajectory log and the agent version. Success is the official AndroidWorld verifier. "
      f"Recordings are not in Git because of their size; every clip is in `{ZIP}` under the run ID. Summary and analysis: "
      "[TEAM_FINAL_EVALUATION.md](TEAM_FINAL_EVALUATION.md).\n"]
for task in ORDER:
    runs = [r for r in RUNS if r[0] == task]
    md.append(f"\n## {task} ({ROW[task]}): {sum(r[5] == 'PASS' for r in runs)}/3\n")
    md.append("| Field | " + " | ".join(f"Run {r[1]} of 3" for r in runs) + " |")
    md.append("| --- | --- | --- | --- |")
    for field in FIELDS:
        cells = []
        for r in runs:
            v = values(r)[field]
            if field == "trajectory":
                v = f"[Readable trajectory log](evidence/team/runs/{r[3]}/trajectory.txt)"
            elif field == "verifier":
                v = f"**{v}**"
            cells.append(str(v).replace("|", "/"))
        md.append(f"| {field} | " + " | ".join(cells) + " |")
(repo / "docs" / "RUN_LOG.md").write_text("\n".join(md) + "\n", encoding="utf-8", newline="\n")

# Word version (landscape, clickable trajectory links to GitHub)
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)   # A4 landscape
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.6))
doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(9.5)
for name in ("Heading 1", "Heading 2"):
    st = doc.styles[name]; st.font.name = "Arial"; st.font.color.rgb = RGBColor(0x07, 0x21, 0x40)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        st.element.rPr.rFonts.attrib.pop(qn(attr), None)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); pr.append(sh)


def link(paragraph, text, url):
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rel)
    r = OxmlElement("w:r"); pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0065BD"); pr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); pr.append(u)
    r.append(pr); t = OxmlElement("w:t"); t.text = text; r.append(t); h.append(r); paragraph._p.append(h)


doc.add_heading("Run log: final evaluation (v65)", 1)
doc.add_paragraph("One table per task; the three runs are the columns. Fields follow the course's run log (D1 p.49), plus a note that "
                  "explains the run, a link to its readable trajectory log on GitHub and the agent version. Success is decided by the "
                  f"official AndroidWorld verifier. Every recording clip is in {ZIP} under the run ID.")
for i, task in enumerate(ORDER):
    runs = [r for r in RUNS if r[0] == task]
    if i:
        doc.add_page_break()
    doc.add_heading(f"{task} ({ROW[task]}): {sum(r[5] == 'PASS' for r in runs)}/3 passed", 2)
    t = doc.add_table(rows=len(FIELDS) + 1, cols=4); t.style = "Table Grid"
    widths = [Cm(3.4), Cm(7.6), Cm(7.6), Cm(7.6)]
    head = ["Field"] + [f"Run {r[1]} of 3" for r in runs]
    for c, txt in enumerate(head):
        cell = t.cell(0, c); cell.text = ""; run = cell.paragraphs[0].add_run(txt); run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(cell, "072140"); cell.width = widths[c]
    for f_i, field in enumerate(FIELDS, start=1):
        label = t.cell(f_i, 0); label.text = ""; lr = label.paragraphs[0].add_run(field); lr.bold = True; shade(label, "E5EFF8")
        label.width = widths[0]
        for c, r in enumerate(runs, start=1):
            cell = t.cell(f_i, c); cell.text = ""; cell.width = widths[c]
            v = values(r)[field]
            if field == "trajectory":
                link(cell.paragraphs[0], "Readable trajectory log", trajectory_url(r[3]))
            else:
                run = cell.paragraphs[0].add_run(str(v))
                if field == "verifier":
                    run.bold = True
                    if v == "FAIL":
                        shade(cell, "FCECE0")
out = HERE / "Run_Log_Final_Evaluation.docx"
doc.save(out)
print("saved", out, "and", repo / "docs" / "RUN_LOG.md")
