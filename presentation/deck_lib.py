"""Small helpers for building an editable 16:9 PowerPoint deck with python-pptx."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

NAVY = RGBColor(0x07, 0x21, 0x40)
BLUE = RGBColor(0x00, 0x65, 0xBD)      # TUM blue: PASS / primary
LIGHT = RGBColor(0x64, 0xA0, 0xC8)
ORANGE = RGBColor(0xE3, 0x72, 0x22)    # failures
GREY = RGBColor(0x5E, 0x67, 0x73)
LINE = RGBColor(0xD9, 0xDE, 0xE5)
PALE = RGBColor(0xF2, 0xF5, 0xF9)
PALE_BLUE = RGBColor(0xE5, 0xEF, 0xF8)
PALE_ORANGE = RGBColor(0xFC, 0xEC, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x2A, 0x35)
FONT = "Arial"
W, H = 13.333, 7.5


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs


def text(slide, x, y, w, h, content, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.1, fill=None, line=None, margin=0.05):
    """Add a text box. content: str (\\n = new paragraph) or list of paragraphs; a paragraph may be a list of
    (text, {bold, color, size}) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is not None:
        box.fill.solid(); box.fill.fore_color.rgb = fill
    if line is not None:
        box.line.color.rgb = line; box.line.width = Pt(1)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin if side in ("top", "bottom") else margin + 0.03))
    paragraphs = content.split("\n") if isinstance(content, str) else content
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for run_text, style in runs:
            r = p.add_run()
            r.text = run_text
            f = r.font
            f.name = FONT
            f.size = Pt(style.get("size", size))
            f.bold = style.get("bold", bold)
            f.italic = style.get("italic", False)
            f.color.rgb = style.get("color", color)
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=6):
    """Bulleted list: each item str or runs list. Uses a dot marker run for simple editing."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        marker = p.add_run(); marker.text = "▪  "; marker.font.name = FONT; marker.font.size = Pt(size); marker.font.color.rgb = BLUE
        runs = [(item, {})] if isinstance(item, str) else item
        for run_text, style in runs:
            r = p.add_run(); r.text = run_text
            r.font.name = FONT; r.font.size = Pt(style.get("size", size)); r.font.bold = style.get("bold", False)
            r.font.color.rgb = style.get("color", color)
    return box


def rect(slide, x, y, w, h, fill=PALE, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(1)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, title, body, accent=BLUE, title_size=16, body_size=13, fill=PALE):
    rect(slide, x, y, w, h, fill=fill)
    rect(slide, x, y, 0.06, h, fill=accent)
    text(slide, x + 0.18, y + 0.1, w - 0.28, 0.4, title, size=title_size, bold=True, color=NAVY)
    if isinstance(body, list):
        bullets(slide, x + 0.18, y + 0.55, w - 0.3, h - 0.6, body, size=body_size, gap=3)
    else:
        text(slide, x + 0.18, y + 0.55, w - 0.3, h - 0.6, body, size=body_size, color=INK)


def header(slide, kicker, title, number, total):
    rect(slide, 0, 0, W, 0.08, fill=BLUE)
    text(slide, 0.6, 0.32, 9, 0.3, kicker.upper(), size=11, bold=True, color=BLUE)
    text(slide, 0.6, 0.58, 12.1, 0.8, title, size=28, bold=True, color=NAVY, spacing=1.0)
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.6), Inches(7.05), Inches(W - 0.6), Inches(7.05)).line.color.rgb = LINE
    text(slide, 0.6, 7.1, 8, 0.3, "Agentic AI in Modern Business  ·  AndroidWorld GUI agent", size=10, color=GREY)
    text(slide, W - 1.6, 7.1, 1.0, 0.3, f"{number}", size=10, color=GREY, align=PP_ALIGN.RIGHT)


def notes(slide, script):
    slide.notes_slide.notes_text_frame.text = script.strip()


def table(slide, x, y, w, h, rows, col_widths=None, size=12, header_fill=NAVY, header_color=WHITE, zebra=True,
          fills=None, bold_first_col=False, align_center_from=None):
    """rows: list of lists of str (first row = header). fills: {(r,c): RGBColor} cell backgrounds."""
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run()
            run.text = str(value)
            run.font.name = FONT
            run.font.size = Pt(size)
            if align_center_from is not None and c >= align_center_from:
                tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
                run.font.bold = True
                run.font.color.rgb = header_color
            else:
                cell.fill.fore_color.rgb = (fills or {}).get((r, c), PALE if zebra and r % 2 == 0 else WHITE)
                run.font.color.rgb = INK
                run.font.bold = bold_first_col and c == 0
            if fills and (r, c) in fills and fills[(r, c)] in (BLUE, ORANGE, NAVY):
                run.font.color.rgb = WHITE
                run.font.bold = True
    return tbl


def bar_chart(slide, x, y, w, h, categories, series, colors=None, horizontal=False, number_format="General",
              legend=True, font_size=12, gap=60, overlap=0, labels=True, axis_max=None, stacked=False, major_unit=None):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    kind = (XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED) if horizontal else \
           (XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED)
    frame = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = frame.chart
    chart.font.name = FONT
    chart.font.size = Pt(font_size)
    chart.has_legend = legend and len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.gap_width = gap
    plot.overlap = overlap if not stacked else 100
    if labels:
        plot.has_data_labels = True
        plot.data_labels.number_format = number_format
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.font.size = Pt(font_size - 1)
        if not stacked:
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(font_size - 1)
    if axis_max is not None:
        va.maximum_scale = axis_max
    if major_unit is not None:
        va.major_unit = major_unit
    if major_unit is not None:
        va.major_unit = major_unit
    va.minimum_scale = 0
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(font_size)
    ca.format.line.color.rgb = LINE
    for i, s in enumerate(plot.series):
        if colors:
            s.format.fill.solid(); s.format.fill.fore_color.rgb = colors[i % len(colors)]
    return chart


def picture(slide, path, x, y, w=None, h=None, border=True):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = LINE; pic.line.width = Pt(1)
    return pic


def movie(slide, path, poster, x, y, w, h):
    return slide.shapes.add_movie(str(path), Inches(x), Inches(y), Inches(w), Inches(h), poster_frame_image=str(poster), mime_type="video/mp4")


def big_number(slide, x, y, w, value, label, color=BLUE, value_size=40, label_size=13):
    text(slide, x, y, w, 0.75, value, size=value_size, bold=True, color=color)
    text(slide, x, y + 0.78, w, 0.7, label, size=label_size, color=GREY)


def arrow(slide, x1, y1, x2, y2, color=GREY, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    line = c.line._get_or_add_ln()
    from lxml import etree
    tail = etree.SubElement(line, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    tail.set("type", "triangle")
    return c
